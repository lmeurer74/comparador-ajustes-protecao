import os
import openpyxl
import io

# Cole aqui todo o dicionário LN_MAP_PADRAO e VARIACOES_DESCRICAO que você já tem
LN_MAP_PADRAO = {
    ("CTS VTS CONFIG.", "CTs Ratios"): "PROT.TCTR1",
    # ... (MANTENHA TODOS OS SEUS MAPEAMENTOS AQUI)
}

VARIACOES_DESCRICAO = {
    ("PROT.RLAU1", "AutoRestore ENBL"): "RestoreEna",
    # ... (MANTENHA TODAS AS SUAS VARIAÇÕES AQUI)
}

def carregar_biblioteca_prioritaria(file_bib_upload, fabricante):
    """Lê o arquivo de upload, ou o arquivo local na pasta, ou gera pela memória"""
    # 1. Prioridade: Arquivo anexado pelo usuário no App
    if file_bib_upload is not None:
        wb = openpyxl.load_workbook(file_bib_upload, data_only=True)
        return _extrair_de_para(wb, fabricante)
    
    # 2. Prioridade: Arquivo local salvo na pasta do projeto
    caminho_local = "biblioteca_de_para.xlsx"
    if os.path.exists(caminho_local):
        wb = openpyxl.load_workbook(caminho_local, data_only=True)
        return _extrair_de_para(wb, fabricante)
    
    # 3. Fallback: Usa os dicionários padrão declarados acima
    return {}, {} # Retornar vazio forçará o motor a usar o LN_MAP_PADRAO dinamicamente

def _extrair_de_para(wb, fabricante):
    ws = wb[fabricante] if fabricante in wb.sheetnames else wb.active
    de_para_cod, de_para_campos = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 5 or not row[4]: continue
        cod_xls, sec, sub, desc, cod_cid = [str(c).strip() for c in row[:5] if c is not None]
        # (Lógica de normalização de chaves que já existe no seu script...)
    return de_para_cod, de_para_campos
