import openpyxl
import re

def limpar_espacos(texto):
    return " ".join(str(texto).strip().split()) if texto is not None else ""

def obter_abas_validas(wb):
    """Filtra apenas abas que ESTÃO VISÍVEIS e são de Grupos/Ajustes"""
    abas_validas = []
    for sheet in wb.worksheets:
        if sheet.sheet_state == 'visible':
            nome = sheet.title.upper()
            if re.search(r'GRUPO\s*[123456]|GROUP\s*[123456]', nome) or "AJUSTE" in nome or "SETTING" in nome:
                abas_validas.append(sheet.title)
    return abas_validas

def extrair_parametros_visiveis(ws, de_para_cod, de_para_campos, dois_por_ln):
    itens_oa = []
    
    # Detecção do modelo (1 ou 2)
    eh_modelo_1 = False
    for r in range(12, min(45, ws.max_row + 1)):
        # IGNORA LINHAS OCULTAS NA DETECÇÃO
        if ws.row_dimensions[r].hidden: continue
        if "." in limpar_espacos(ws.cell(r, 2).value):
            eh_modelo_1 = True
            break

    # Loop de extração de parâmetros
    start_row = 14 if eh_modelo_1 else 12
    secao, subgrupo = "", ""

    for r in range(start_row, ws.max_row + 1):
        # REGRA CRÍTICA: Ignorar linhas ocultas no Excel!
        if ws.row_dimensions[r].hidden:
            continue
            
        b_val = limpar_espacos(ws.cell(r, 2).value)
        c_val = limpar_espacos(ws.cell(r, 3).value)
        
        if b_val.upper() in ["PARAMETRO", "PARÂMETRO"]: continue
        if b_val.upper() == "ELABORAÇÃO E IMPLANTAÇÃO": break
        
        # (COLE AQUI A LÓGICA DE EXTRAÇÃO MODELO 1 / MODELO 2 QUE VOCÊ JÁ TEM)
        # item = {"col_a": ..., "col_b": ..., "col_c": ..., "cod_cid": ...}
        # itens_oa.append(item)

    return itens_oa
